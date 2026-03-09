import io
import logging
from pathlib import Path
from typing import List

from pptx import Presentation
from PIL import Image

from private_gpt.components.ocr_components.discovery.base import (
    BaseDiscoveryProvider,
    DiscoveryElement,
    DiscoveryResult,
    DiscoveryType,
)

logger = logging.getLogger(__name__)

class PPTXDiscoveryProvider(BaseDiscoveryProvider):
    """
    Discovery Provider for Microsoft PowerPoint (.pptx) files.
    
    Scans slide shapes to extract digital text and image objects.
    """

    @property
    def supported_extensions(self) -> List[str]:
        return ["pptx"]

    def discover(self, file_path: Path) -> DiscoveryResult:
        logger.info("Opening PowerPoint for discovery: %s", file_path.name)
        prs = Presentation(file_path)
        elements = []

        for i, slide in enumerate(prs.slides):
            page_num = i + 1
            for shape in slide.shapes:
                # 1. Extract Text from shapes (TextFrame)
                if hasattr(shape, "text_frame") and shape.text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        elements.append(DiscoveryElement(
                            element_type=DiscoveryType.TEXT,
                            content=text,
                            page_number=page_num,
                            metadata={"shape_name": shape.name, "type": "slide_text"}
                        ))

                # 2. Extract Images (Picture)
                if shape.shape_type == 13: # 13 is the PICTURE shape type
                    try:
                        image_bytes = shape.image.blob
                        img = Image.open(io.BytesIO(image_bytes)).convert("RGB")
                        
                        elements.append(DiscoveryElement(
                            element_type=DiscoveryType.IMAGE,
                            content=img,
                            page_number=page_num,
                            metadata={
                                "shape_name": shape.name,
                                "type": "slide_image",
                                "ext": shape.image.ext
                            }
                        ))
                        logger.debug("Discovered image on slide %d: %s", page_num, shape.name)
                    except Exception as e:
                        logger.warning("Failed to extract image on slide %d: %s", page_num, str(e))

        logger.info("Discovery complete for %s. Found %d elements.", file_path.name, len(elements))
        return DiscoveryResult(
            file_path=file_path,
            file_type="pptx",
            elements=elements
        )
